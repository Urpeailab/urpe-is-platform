"""
Expert Letters Router — Professional expert opinion letters for NIW petitions
Extracted from server.py for better code organization.
"""

import io
import re
import json
import uuid
import logging
from datetime import datetime, timezone
from typing import Optional

from fastapi import APIRouter, BackgroundTasks, Depends, File, Form, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from fastapi.security import HTTPBearer, HTTPAuthorizationCredentials

import sys, os
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from pdf_utils import pdf_safe as _pdf_safe

from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.lib.enums import TA_LEFT, TA_JUSTIFY, TA_CENTER
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer




def _pick_profile_for_signer(title: str, credentials: str, organization: str) -> dict:
    """
    Choose a writing-style profile for the letter.

    Uses the 60/40 weighted picker from letter_format_profiles: 60% of the
    time we pick the profile suggested by the signer's keywords (so
    academic-looking signers tend to get an academic voice), 40% of the
    time we pick a completely random profile so two letters signed by the
    same person don't look identical. This is the main lever that breaks
    the "all letters look the same" complaint.
    """
    from letter_format_profiles import pick_profile_for_signer
    signer_blob = f"{title} {credentials} {organization}"
    return pick_profile_for_signer(signer_blob, random_ratio=0.4)


def _clean_placeholders(text: str, today_str: str) -> str:
    """
    Aggressively remove ALL [PLACEHOLDER] patterns from generated letters.
    In immigration letters, square brackets should NEVER appear in final output.
    """
    if not text:
        return text

    # 1. Strip markdown code fences — both plain AND wrapped in <p> tags
    text = text.strip()
    text = re.sub(r'^```(?:html|markdown)?\s*\n?', '', text)
    text = re.sub(r'\n?```\s*$', '', text.strip())
    # Handle code fences wrapped in <p> tags by Gemini: <p>```html</p> or <p>```</p>
    text = re.sub(r'<p>\s*```+\s*(?:html|markdown)?\s*</p>\s*', '', text, flags=re.IGNORECASE)
    text = re.sub(r'<p>\s*```+\s*</p>\s*', '', text, flags=re.IGNORECASE)

    # 2. Replace date-specific placeholders with today's date
    date_patterns = [
        r'\[DATE\]', r'\[TODAY\]', r'\[CURRENT\s+DATE\]', r'\[CURRENT_DATE\]',
        r'\[Month\s+DD,?\s*YYYY\]', r'\[Month\s+\d+,?\s*\d{4}\]',
        r'\[Month,?\s*Year\]', r'\[MM/DD/YYYY\]',
    ]
    for p in date_patterns:
        text = re.sub(p, today_str, text, flags=re.IGNORECASE)

    # 3. Remove entire lines that consist solely of a placeholder
    lines = text.split('\n')
    lines = [ln for ln in lines if not re.match(r'^\s*\[[^\[\]]+\]\s*$', ln)]
    text = '\n'.join(lines)

    # 4. Remove ALL remaining [bracket content] inline — any length, any case
    #    Keep only [1], [2] single-digit citation refs (but immigration letters rarely use these)
    text = re.sub(r'\[[^\[\]]+\]', lambda m: m.group(0) if re.match(r'^\[\d\]$', m.group(0)) else '', text)

    # 5. Remove empty <p> tags left after bracket removal
    text = re.sub(r'<p>\s*</p>', '', text)
    text = re.sub(r'<p>\s*[,.:;]\s*</p>', '', text)

    # 6. Fix sentence artifacts from removed placeholders
    text = re.sub(r'  +', ' ', text)                        # collapse double spaces
    text = re.sub(r' ([,.:;])', r'\1', text)                # space before punctuation
    text = re.sub(r',\s*,', ',', text)                      # double commas
    text = re.sub(r'\(\s*\)', '', text)                     # empty parentheses
    text = re.sub(r'\n{3,}', '\n\n', text)                  # max 2 blank lines

    # 7. Fix common awkward phrases left after placeholder removal
    text = re.sub(r'including\s*[,.]', 'including relevant professional associations.', text, flags=re.IGNORECASE)
    text = re.sub(r'member of\s*[,.]', 'member of several professional organizations.', text, flags=re.IGNORECASE)
    text = re.sub(r'such as\s*[,.]', 'such as relevant organizations in the field.', text, flags=re.IGNORECASE)

    return text

logger = logging.getLogger(__name__)

# ──────────────────────────────────────────────
# Module-level dependencies (injected via init_router)
# ──────────────────────────────────────────────
_db = None
_get_current_user = None
_call_openai_gpt4o = None
_call_gemini_flash_lite = None


def init_router(
    database,
    get_current_user_func,
    call_openai_gpt4o_func,
    call_gemini_flash_lite_func,
):
    global _db, _get_current_user, _call_openai_gpt4o, _call_gemini_flash_lite
    _db = database
    _get_current_user = get_current_user_func
    _call_openai_gpt4o = call_openai_gpt4o_func
    _call_gemini_flash_lite = call_gemini_flash_lite_func
    logger.info("✅ Expert Letters Router initialized with dependencies")


def get_db():
    if _db is None:
        raise RuntimeError("Database not initialized. Call init_router() first.")
    return _db


async def get_current_user_wrapper(credentials: HTTPAuthorizationCredentials = Depends(HTTPBearer())):
    if _get_current_user is None:
        raise RuntimeError("get_current_user not initialized.")
    return await _get_current_user(credentials)


# ──────────────────────────────────────────────
# Text extraction helper (inline — only used by expert letters)
# ──────────────────────────────────────────────

async def _extract_text_from_upload(file: UploadFile) -> str:
    """Extract text from an uploaded file (PDF, DOCX, or plain text)."""
    content = await file.read()
    fname = (file.filename or "").lower()

    if fname.endswith('.pdf'):
        try:
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(content))
            return "\n".join(page.extract_text() or '' for page in reader.pages)
        except ImportError:
            try:
                from PyPDF2 import PdfReader
                reader = PdfReader(io.BytesIO(content))
                return "\n".join(page.extract_text() or '' for page in reader.pages)
            except Exception:
                pass
    elif fname.endswith('.docx'):
        try:
            import docx as python_docx
            doc = python_docx.Document(io.BytesIO(content))
            return "\n".join(p.text for p in doc.paragraphs)
        except Exception:
            pass

    return content.decode('utf-8', errors='ignore')


def _generate_expert_letter_pdf(
    letter_text: str,
    client_name: str,
    expert_name: str,
    language: str,
    profile: dict = None,
) -> bytes:
    """Generate PDF from expert letter markdown content using the given format profile."""
    from letter_format_profiles import get_profile, resolve_fonts
    if profile is None:
        profile = get_profile("classic_legal")
    p = dict(profile["pdf"])   # shallow copy so we can splice in resolved fonts
    # Resolve font_intent (new system) into concrete font names usable by
    # ReportLab. The renderer below still reads `font_body` / `font_bold` /
    # `font_italic`, so we expose them here to keep that interface intact.
    _fonts = resolve_fonts(profile)
    p["font_body"] = _fonts["regular"]
    p["font_bold"] = _fonts["bold"]
    p["font_italic"] = _fonts["italic"]

    buffer = io.BytesIO()

    # Clean up markdown artifacts
    letter_text = re.sub(r'<b>([^<]*)<b>', r'<b>\1</b>', letter_text)
    letter_text = re.sub(r'<i>([^<]*)<i>', r'<i>\1</i>', letter_text)

    # Convertir markdown inline (**bold**, *italic*, runs de 3+ asteriscos) a XML
    # inline de ReportLab. Aplica también dentro de tags HTML como <p>**X**</p>.
    from pdf_utils import md_inline_to_rl
    letter_text = _pdf_safe(letter_text)
    letter_text = md_inline_to_rl(letter_text)

    from reportlab.lib.units import inch as _inch
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        topMargin=p["top_margin"],
        bottomMargin=p["bottom_margin"],
        leftMargin=p["left_margin"],
        rightMargin=p["right_margin"],
        title=f"Expert Opinion Letter — {client_name}",
        author=expert_name,
        subject=f"EB-2 NIW Expert Opinion Letter",
        creator="Monica NIW System",
    )

    styles = getSampleStyleSheet()
    section_style = p["section_style"]

    # Build section title style based on profile
    if section_style == "ALL_CAPS_UNDERLINE":
        h1_style = ParagraphStyle(
            'H1Custom', fontSize=p["font_size_section"], fontName=p["font_bold"],
            textColor=colors.black, spaceAfter=6, spaceBefore=p["space_before_section"],
            underline=True, alignment=TA_LEFT
        )
    elif section_style == "BOLD_RULE":
        h1_style = ParagraphStyle(
            'H1Custom', fontSize=p["font_size_section"] + 1, fontName=p["font_bold"],
            textColor=colors.HexColor('#1a1a2e'), spaceAfter=2, spaceBefore=p["space_before_section"],
            borderPadding=(0, 0, 4, 0), alignment=TA_LEFT
        )
    elif section_style == "ROMAN_NUMERAL":
        h1_style = ParagraphStyle(
            'H1Custom', fontSize=p["font_size_section"] + 1, fontName=p["font_bold"],
            textColor=colors.black, spaceAfter=8, spaceBefore=p["space_before_section"],
            alignment=TA_LEFT
        )
    elif section_style == "BOLD_LEFT_ACCENT":
        h1_style = ParagraphStyle(
            'H1Custom', fontSize=p["font_size_section"], fontName=p["font_bold"],
            textColor=colors.HexColor('#2d3748'), spaceAfter=6, spaceBefore=p["space_before_section"],
            leftIndent=8, borderPadding=(0, 0, 0, 4), alignment=TA_LEFT
        )
    elif section_style == "NUMBERED_BOLD":
        h1_style = ParagraphStyle(
            'H1Custom', fontSize=p["font_size_section"], fontName=p["font_bold"],
            textColor=colors.black, spaceAfter=6, spaceBefore=p["space_before_section"],
            alignment=TA_LEFT
        )
    else:
        h1_style = ParagraphStyle(
            'H1Custom', fontSize=p["font_size_section"], fontName=p["font_bold"],
            textColor=colors.black, spaceAfter=6, spaceBefore=p["space_before_section"],
        )

    body_style = ParagraphStyle(
        'BodyCustom', fontSize=p["font_size_body"], leading=p["leading"],
        alignment=TA_JUSTIFY, spaceAfter=p["space_after_para"],
        fontName=p["font_body"], textColor=colors.black
    )
    meta_style = ParagraphStyle(
        'MetaCustom', fontSize=p["font_size_body"] - 1, leading=p["leading"] - 3,
        alignment=TA_LEFT, spaceAfter=4, fontName=p["font_body"], textColor=colors.HexColor('#555555')
    )

    content = []

    # ── Clean content before PDF generation ────────────────────────────────────
    letter_text = letter_text.strip()
    letter_text = re.sub(r'^```(?:html|markdown)?\s*\n?', '', letter_text)
    letter_text = re.sub(r'\n?```\s*$', '', letter_text.strip())
    letter_text = re.sub(r'<p>\s*`{1,3}\s*(?:html|markdown)?\s*</p>\s*', '', letter_text, flags=re.IGNORECASE)
    letter_text = re.sub(r'\[[^\[\]]+\]', lambda m: m.group(0) if re.match(r'^\[\d\]$', m.group(0)) else '', letter_text)
    letter_text = re.sub(r'<p>\s*</p>', '', letter_text)
    letter_text = re.sub(r'  +', ' ', letter_text)
    letter_text = re.sub(r' ([,.:;])', r'\1', letter_text)

    # ── HTML-aware content builder ──────────────────────────────────────────────
    from bs4 import BeautifulSoup as _BS

    def _rl_inline(tag):
        parts = []
        for child in tag.children:
            if hasattr(child, 'name') and child.name:
                txt = child.get_text()
                if child.name in ('strong', 'b'):
                    parts.append(f'<b>{txt}</b>')
                elif child.name in ('em', 'i'):
                    parts.append(f'<i>{txt}</i>')
                else:
                    parts.append(txt)
            else:
                parts.append(str(child))
        return ''.join(parts).strip()

    if '<p' in letter_text or '<div' in letter_text:
        # HTML path — parse with BeautifulSoup
        soup = _BS(letter_text, 'html.parser')
        for tag in soup.find_all(['p', 'h1', 'h2', 'h3', 'h4', 'div']):
            if tag.name in ('h1', 'h2', 'h3', 'h4'):
                txt = tag.get_text().strip()
                if txt:
                    content.append(Paragraph(_pdf_safe(txt), h1_style))
                    if section_style == "BOLD_RULE":
                        from reportlab.platypus import HRFlowable
                        content.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor('#cccccc'), spaceAfter=6))
            elif tag.name in ('p', 'div'):
                if tag.find(['p', 'div']):
                    continue  # skip container divs
                inline = _rl_inline(tag)
                if inline:
                    try:
                        content.append(Paragraph(inline, body_style))
                    except Exception:
                        content.append(Paragraph(_pdf_safe(tag.get_text().strip()), body_style))
    else:
        # Plain-text / Markdown path (legacy fallback)
        lines = letter_text.split('\n')
        current_paragraph = []

        def flush_para():
            if current_paragraph:
                content.append(Paragraph(' '.join(current_paragraph), body_style))
                current_paragraph.clear()

        for line in lines:
            line = line.strip()
            if not line:
                flush_para()
                continue
            if re.match(r'^\[.*\]$', line):
                continue
            is_section = (
                (line.isupper() and 2 <= len(line.split()) <= 10)
                or re.match(r'^[IVX]+\.\s+[A-Z]', line)
                or re.match(r'^\d+\.\s+[A-Z][A-Z ]', line)
                or re.match(r'^#{1,3}\s+', line)
            )
            if is_section:
                flush_para()
                clean_line = re.sub(r'^#{1,3}\s+', '', line)
                content.append(Paragraph(_pdf_safe(clean_line), h1_style))
                if section_style == "BOLD_RULE":
                    from reportlab.platypus import HRFlowable
                    content.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor('#cccccc'), spaceAfter=6))
                continue
            if line.startswith(('───', '---', '===')) and len(line) > 4:
                flush_para()
                content.append(Spacer(1, 0.1 * _inch))
                continue
            if line.startswith(('• ', '- ', '* ', '· ')):
                flush_para()
                content.append(Paragraph('• ' + _pdf_safe(line[2:]), ParagraphStyle(
                    'Bullet', parent=body_style, leftIndent=20, firstLineIndent=-12, spaceAfter=4
                )))
                continue
            if re.match(r'^(RE:|Re:|Ref:|SUBJECT:|Sincerely|Respectfully|Cordially|Dear|Estimad|Atentamente|Warmly|Faithfully|Yours|With sincere|With warm|With great|Best regards|Kind regards|Regards|Most respectfully)', line):
                flush_para()
                content.append(Paragraph(_pdf_safe(line), body_style))
                # Leave ~1 inch of vertical whitespace after the sign-off so
                # the signer can sign by hand. Without this the printed name
                # sits right under "Sincerely," and there is no room for a
                # physical signature.
                content.append(Spacer(1, 70))
                continue
            current_paragraph.append(_pdf_safe(line))

        flush_para()

    # Post-pass safety net: if the HTML path produced the sign-off as a
    # <p> tag, we never hit the plain-text branch above. Walk `content` and
    # ensure there's ~70pt of whitespace after the first sign-off paragraph.
    _SIGN_OFF_RX = re.compile(
        r'^\s*(Sincerely|Respectfully|Cordially|Warmly|Faithfully|Yours|'
        r'With sincere|With warm|With great|Best regards|Kind regards|'
        r'Regards|Most respectfully)[,\.]?\s*$',
        re.IGNORECASE,
    )
    for idx in range(len(content) - 1):
        flow = content[idx]
        if isinstance(flow, Paragraph):
            text = re.sub(r'<[^>]+>', '', flow.text or '').strip()
            if _SIGN_OFF_RX.match(text):
                # Already has a tall spacer? Skip.
                nxt = content[idx + 1] if idx + 1 < len(content) else None
                if isinstance(nxt, Spacer) and getattr(nxt, 'height', 0) >= 60:
                    break
                content.insert(idx + 1, Spacer(1, 70))
                break

    doc.build(content)
    buffer.seek(0)
    return buffer.getvalue()


# ──────────────────────────────────────────────
# Router and endpoints
# ──────────────────────────────────────────────

router = APIRouter(prefix="/expert-letters", tags=["Expert Letters"])


@router.post("/generate")
async def generate_expert_letter(
    background_tasks: BackgroundTasks,
    client_cv: UploadFile = File(..., description="CV del Cliente (obligatorio)"),
    project_info: UploadFile = File(..., description="Información del Proyecto (obligatorio)"),
    expert_cv: UploadFile = File(..., description="CV del Experto/Firmante (obligatorio)"),
    client_id: Optional[str] = Form(None),
    current_user=Depends(get_current_user_wrapper),
):
    """Start async expert letter generation. Returns immediately with letter_id and status='generating'."""
    db = get_db()
    letter_id = str(uuid.uuid4())

    logger.info(f"📝 Starting expert letter generation | client_id={client_id}")

    # Read file contents NOW — UploadFile cannot be used after request context closes
    client_cv_text = await _extract_text_from_upload(client_cv)
    project_info_text = await _extract_text_from_upload(project_info)
    expert_cv_text = await _extract_text_from_upload(expert_cv)

    # Create letter document immediately with generating status
    from letter_format_profiles import pick_random_profile
    format_profile = pick_random_profile()
    logger.info(f"🎨 Using format profile: {format_profile['name']}")

    letter_doc = {
        "id": letter_id,
        "user_id": current_user.id,
        "client_id": client_id,
        "client_name": "",
        "expert_name": "",
        "expert_organization": "",
        "project_title": "",
        "content_en": "",
        "content_es": "",
        "current_language": "en",
        "status": "generating",
        "format_profile_id": format_profile["id"],
        "progress_percentage": 5,
        "progress_message": "Extrayendo texto de los documentos...",
        "created_at": datetime.now(timezone.utc).isoformat(),
        "updated_at": datetime.now(timezone.utc).isoformat()
    }
    await db.expert_letters.insert_one(letter_doc)
    logger.info(f"💾 Expert letter record created: {letter_id}")

    background_tasks.add_task(
        _generate_expert_letter_background,
        letter_id, client_cv_text, project_info_text, expert_cv_text, format_profile
    )

    return {"letter_id": letter_id, "status": "generating", "message": "Generación iniciada en segundo plano"}


async def _generate_expert_letter_background(
    letter_id: str,
    client_cv_text: str,
    project_info_text: str,
    expert_cv_text: str,
    format_profile: dict,
):
    """Background task: runs the LLM calls and updates the letter document."""
    db = get_db()
    logger.info(f"🚀 Background generation started for letter {letter_id}")

    async def _update_progress(pct: int, msg: str):
        await db.expert_letters.update_one(
            {"id": letter_id},
            {"$set": {"progress_percentage": pct, "progress_message": msg, "updated_at": datetime.now(timezone.utc).isoformat()}}
        )

    try:
        from expert_letter_endpoints import EXPERT_LETTER_SYSTEM_PROMPT, ANTI_PLACEHOLDER_RULE
        from datetime import datetime as dt
        today_str = dt.now().strftime("%B %d, %Y")

        await _update_progress(15, "Analizando documentos con IA...")

        analysis_prompt = f"""You are analyzing three documents for generating a professional EB-2 NIW expert opinion letter aligned with Matter of Dhanasar, 26 I&N Dec. 884 (AAO 2016).

**DOCUMENT 1 - CLIENT'S CV:**
{client_cv_text[:8000]}

**DOCUMENT 2 - PROJECT INFORMATION:**
{project_info_text[:8000]}

**DOCUMENT 3 - EXPERT'S CV:**
{expert_cv_text[:6000]}

Extract and return ONLY a JSON object with this structure:
{{
  "client_name": "Full name from client CV",
  "client_field": "Specific field of expertise (e.g., 'Healthcare AI and Clinical Decision Support', not just 'technology')",
  "client_position": "Current position and organization",
  "client_achievements": "Key publications, patents, awards, specific accomplishments with numbers",
  "client_unique_skills": "3-4 unique skills or skill combinations that are rare in the U.S. workforce",
  "project_title": "Title/name of the project",
  "project_description": "Detailed description of the project and how it works",
  "project_importance": "Why this project matters at the NATIONAL level (not just locally)",
  "project_technical_details": "Technical specifics: methodologies, technologies, innovations",
  "project_geographic_reach": "States, organizations, or people affected nationally",
  "project_economic_impact": "Economic value, cost savings, revenue, or jobs created/supported",
  "project_federal_alignment": "Which federal laws or programs this project aligns with",
  "expert_name": "Full name from expert CV",
  "expert_title": "Professional title/position",
  "expert_organization": "Current institution/organization",
  "expert_credentials": "Key qualifications, publications, years of experience, awards",
  "expert_email": "Email if available",
  "expert_phone": "Phone if available"
}}

IMPORTANT: Extract real, specific information from all three documents. Be thorough and precise."""

        analysis_raw = await _call_openai_gpt4o(
            "You are an expert at analyzing documents and extracting structured information. Return only valid JSON.",
            analysis_prompt,
            temperature=0.3,
            max_tokens=2000
        )

        analysis_raw = analysis_raw.strip()
        if analysis_raw.startswith('```'):
            analysis_raw = analysis_raw.split('```')[1]
            if analysis_raw.startswith('json'):
                analysis_raw = analysis_raw[4:]
        extracted_data = json.loads(analysis_raw.strip())
        logger.info(f"✅ Document analysis complete for {letter_id}")

        # ── Select profile based on signer's credentials ──────────────────────
        # 60/40 weighted picker: usually the suggested voice for the signer's
        # credentials, but 40% of the time a random voice so two letters from
        # the same signer don't read identically.
        format_profile = _pick_profile_for_signer(
            extracted_data.get('expert_title', ''),
            extracted_data.get('expert_credentials', ''),
            extracted_data.get('expert_organization', '')
        )

        # Pick a sign-off compatible with this voice's mood and a "style salt"
        # to break LLM caching. These vary on every call.
        from letter_format_profiles import (
            pick_sign_off as _pick_sign_off,
            make_style_salt as _make_style_salt,
            pick_temperature as _pick_temperature,
        )
        chosen_sign_off = _pick_sign_off(format_profile)
        style_salt = _make_style_salt()
        gen_temperature = _pick_temperature()

        await _update_progress(40, f"Redactando carta en inglés para {extracted_data.get('client_name', 'el cliente')}...")

        generation_prompt = f"""Generate a professional expert opinion letter following the Dhanasar 3-Prong NIW structure.

{style_salt}

**CURRENT DATE: {today_str}** — Use this exact date. Never invent a different date.

{format_profile['prompt_instructions']}

---
## CANDIDATE (PETITIONER) INFORMATION:
- Full Name: {extracted_data['client_name']}
- Field of Expertise: {extracted_data['client_field']}
- Current Position: {extracted_data['client_position']}
- Key Achievements & Credentials: {extracted_data['client_achievements']}

## PROJECT / PROPOSED ENDEAVOR:
- Title: {extracted_data.get('project_title', 'N/A')}
- Description: {extracted_data.get('project_description', 'N/A')}
- National Importance: {extracted_data.get('project_importance', 'N/A')}
- Technical Details: {extracted_data.get('project_technical_details', 'N/A')}

## EXPERT (LETTER SIGNER) INFORMATION:
- Full Name: {extracted_data['expert_name']}
- Title: {extracted_data['expert_title']}
- Organization: {extracted_data['expert_organization']}
- Credentials & Expertise: {extracted_data['expert_credentials']}
- Email: {extracted_data.get('expert_email', '')}
- Phone: {extracted_data.get('expert_phone', '')}

---
## MANDATORY FORMAT — Flowing Professional Letter (NO section headers):

CRITICAL FORMAT RULES:
- Write as a continuous professional letter with 12-14 flowing paragraphs.
- DO NOT use section titles, section numbers, or heading-style text of any kind.
- DO NOT label sections as "SECTION 1", "PRONG 1:", "I.", "II.", "CONCLUSION:", or similar.
- DO NOT use <h1>, <h2>, <h3>, <h4>, or any HTML heading tags.
- Use ONLY <p> tags for all paragraph content.
- Use <strong> only for metrics or key terms WITHIN a paragraph, never as a standalone title.
- Do NOT insert empty <p> or <br> tags between paragraphs.
- The letter reads as a traditional professional narrative — identical in format to the reference letter.

LETTER CONTENT (one paragraph per item — no titles, no labels):

Paragraph 1 (Header + Opening): Write date ({today_str}), USCIS address block, and "Re:" line as plain text in the paragraph. Salutation: "Dear USCIS Adjudicating Officer:". Then expert introduces themselves: full name, title, institution, years of experience. State purpose: evaluating {extracted_data['client_name']}'s EB-2 NIW petition. Invoke Matter of Dhanasar, 26 I&N Dec. 884 (AAO 2016).

Paragraph 2 (Expert Authority): Expert's key publications, grants, awards. Establishes credibility and why uniquely qualified to evaluate this specific field.

Paragraph 3 (Independence + Documents Reviewed): MANDATORY: "I have no prior personal or professional relationship with {extracted_data['client_name']}. My evaluation is based solely on an independent review of [his/her] documented body of work, including [list 2-3 specific documents reviewed]." State that this independence allows for an objective assessment.

Paragraph 4 (National Gap): Cite a REAL government statistic (BLS, NIH, NSF, DOE, USDA, HUD, DOT, SBA, CMS, or EPA) to document the national gap or workforce shortage in the field "{extracted_data['client_field']}". State the cost of not addressing it.

Paragraph 5 (Proposed Endeavor + Substantial Merit): Describe the proposed endeavor in specific technical terms. What it is, how it works, why it has substantial merit. No heading — just flows from the national gap paragraph.

Paragraph 6 (National Importance + Federal Law): Why this work matters at the NATIONAL level. Connect to specific U.S. federal laws or programs (full names with H.R./P.L. numbers). Geographic reach: states, organizations, beneficiaries. Close this paragraph with expert opinion that the endeavor satisfies Prong 1 of Matter of Dhanasar.

Paragraphs 7-9 (Three Achievement Paragraphs — Well Positioned): Each paragraph presents one specific achievement of {extracted_data['client_name']} that the expert verified. Format within each paragraph: what was accomplished + [baseline] → [result] ([X%] change) over [period] across [N cases] + how the expert verified it.

Paragraph 10 (Unique Skill Combination + Prong 2 Conclusion): Why {extracted_data['client_name']}'s specific combination of skills is exceptional and rare in the U.S. workforce. Close with expert opinion that Prong 2 of Matter of Dhanasar is satisfied.

Paragraph 11 (Time-Sensitive National Need): The national need is urgent. Labor certification delays of 12-18 months would cost the U.S. [specific consequence]. The standard process cannot identify a U.S. worker with this exact combination of skills.

Paragraph 12 (Economic Multiplier + Prong 3 Conclusion): Quantify economic value: direct impact ($), jobs supported, beneficiaries, 3-year projection. Close with expert opinion that Prong 3 of Matter of Dhanasar is satisfied.

Paragraph 13 (Comparative Exceptionality): In the expert's years evaluating professionals in this field, what makes {extracted_data['client_name']} stand out among peers. Why the combination is rare.

Paragraph 14 (Conclusion + Strongest Recommendation): Summarize satisfaction of all 3 Dhanasar prongs in flowing prose (not a checklist). Provide strongest expert recommendation for the NIW petition. State availability for follow-up.

After paragraph 14, close with:
"{chosen_sign_off}"
[Expert's full name, title, organization, email, phone]

---
## CRITICAL REQUIREMENTS:
1. DO NOT invent specific statistics — use real government agency ranges for the relevant field.
2. All metrics MUST include: baseline + result + % change + time period + sample size + how expert verified.
3. Cite "Matter of Dhanasar, 26 I&N Dec. 884 (AAO 2016)" explicitly in the letter.
4. Independence clause MANDATORY.
5. Use {today_str} as the letter date.
6. Replace ALL placeholder text — ZERO brackets in output.
7. Length: 2,500-3,500 words.
8. OUTPUT FORMAT: ONLY <p> tags. NEVER use <h1>, <h2>, <h3>, <h4>. No HTML heading tags whatsoever."""

        from llm_fallback import call_llm_with_fallback
        letter_content_en = await call_llm_with_fallback(
            system_prompt=EXPERT_LETTER_SYSTEM_PROMPT + ANTI_PLACEHOLDER_RULE,
            user_prompt=generation_prompt,
            primary_gemini_fn=_call_gemini_flash_lite,
            temperature=gen_temperature,
            max_tokens=8000,
            min_chars=500,
            label=f"expert-letter/{letter_id[:8]}",
        )
        logger.info(
            f"🎨 expert letter generated | profile={format_profile['id']} "
            f"sign_off={chosen_sign_off!r} temp={gen_temperature}"
        )

        await _update_progress(75, "Traduciendo carta al español...")
        logger.info(f"✅ English letter generated for {letter_id}, translating to Spanish...")

        translation_prompt = f"""Translate this professional expert opinion letter from English to Spanish.

IMPORTANT TRANSLATION RULES:
1. Maintain the professional tone
2. Keep all proper nouns in English
3. Keep technical terms accurate
4. Preserve structure and formatting
5. Use formal Spanish (usted form)

ENGLISH LETTER:
{letter_content_en}

Provide ONLY the Spanish translation, maintaining exact structure."""

        letter_content_es = await _call_openai_gpt4o(
            "You are an expert legal translator. Translate with precision and maintain professional tone.",
            translation_prompt,
            temperature=0.2,
            max_tokens=6000
        )

        await _update_progress(95, "Guardando carta en el sistema...")

        letter_content_en = _clean_placeholders(letter_content_en, today_str)
        letter_content_es = _clean_placeholders(letter_content_es, today_str)

        await db.expert_letters.update_one(
            {"id": letter_id},
            {"$set": {
                "client_name": extracted_data['client_name'],
                "client_field": extracted_data['client_field'],
                "expert_name": extracted_data['expert_name'],
                "expert_organization": extracted_data['expert_organization'],
                "project_title": extracted_data.get('project_title', ''),
                "content_en": letter_content_en,
                "content_es": letter_content_es,
                "current_language": "en",
                "status": "completed",
                "progress_percentage": 100,
                "progress_message": "Carta generada exitosamente",
                "updated_at": datetime.now(timezone.utc).isoformat()
            }}
        )
        logger.info(f"✅ Expert letter {letter_id} completed and saved")

    except Exception as e:
        logger.error(f"❌ Background generation failed for {letter_id}: {str(e)}")
        import traceback
        logger.error(traceback.format_exc())
        await db.expert_letters.update_one(
            {"id": letter_id},
            {"$set": {
                "status": "error",
                "progress_message": f"Error: {str(e)}",
                "error_message": str(e),
                "updated_at": datetime.now(timezone.utc).isoformat()
            }}
        )


@router.get("")
async def list_expert_letters(
    client_id: Optional[str] = Query(None),
    current_user=Depends(get_current_user_wrapper),
):
    """List all expert letters (excluding deleted)."""
    db = get_db()
    query = {"status": {"$ne": "deleted"}}
    if client_id:
        query["client_id"] = client_id

    letters = await db.expert_letters.find(query, {"_id": 0}).to_list(1000)
    logger.info(f"📋 Retrieved {len(letters)} expert letters")
    return {"letters": letters}


@router.get("/{letter_id}")
async def get_expert_letter(letter_id: str, current_user=Depends(get_current_user_wrapper)):
    """Get specific expert letter — content cleaned on-the-fly."""
    db = get_db()
    letter = await db.expert_letters.find_one({"id": letter_id}, {"_id": 0})
    if not letter:
        raise HTTPException(status_code=404, detail="Letter not found")
    # Clean content retroactively for letters stored with old format
    if letter.get('content_en'):
        letter['content_en'] = _clean_placeholders(letter['content_en'], datetime.now().strftime("%B %d, %Y"))
    if letter.get('content_es'):
        letter['content_es'] = _clean_placeholders(letter['content_es'], datetime.now().strftime("%d de %B de %Y"))
    return letter


@router.get("/{letter_id}/download")
async def download_expert_letter_pdf(
    letter_id: str,
    language: str = Query("en"),
    current_user=Depends(get_current_user_wrapper),
):
    """Download expert letter as PDF."""
    db = get_db()
    letter_doc = await db.expert_letters.find_one({"id": letter_id}, {"_id": 0})
    if not letter_doc:
        raise HTTPException(status_code=404, detail="Letter not found")

    letter_text = letter_doc.get('content_es' if language == 'es' else 'content_en')
    if not letter_text:
        raise HTTPException(status_code=404, detail=f"Letter content not available in {language}")

    client_name = letter_doc.get('client_name', 'Client')
    expert_name = letter_doc.get('expert_name', 'Expert')

    logger.info(f"📥 Generating PDF for letter {letter_id} in {language}")
    from letter_format_profiles import get_profile
    profile = get_profile(letter_doc.get('format_profile_id', 'classic_legal'))
    pdf_bytes = _generate_expert_letter_pdf(letter_text, client_name, expert_name, language, profile)

    client_clean = client_name.replace(' ', '_')
    expert_clean = expert_name.replace(' ', '_')
    lang_suffix = "ES" if language == "es" else "EN"
    filename = f"Carta_Experto_{client_clean}_firmada_por_{expert_clean}_{lang_suffix}.pdf"

    logger.info(f"✅ PDF generated: {filename}")
    return StreamingResponse(
        io.BytesIO(pdf_bytes),
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


@router.get("/{letter_id}/download-docx")
async def download_expert_letter_docx(
    letter_id: str,
    language: str = Query("en"),
    current_user=Depends(get_current_user_wrapper),
):
    """Download expert letter as Microsoft Word (.docx) — opens cleanly in Google Docs."""
    from docx_utils import build_docx_response

    db = get_db()
    letter_doc = await db.expert_letters.find_one({"id": letter_id}, {"_id": 0})
    if not letter_doc:
        raise HTTPException(status_code=404, detail="Letter not found")
    letter_text = letter_doc.get('content_es' if language == 'es' else 'content_en')
    if not letter_text:
        raise HTTPException(status_code=404, detail=f"Letter content not available in {language}")

    client_name = letter_doc.get('client_name', 'Client')
    expert_name = letter_doc.get('expert_name', 'Expert')
    client_clean = client_name.replace(' ', '_')
    expert_clean = expert_name.replace(' ', '_')

    is_html = bool(re.search(r'<(p|h[1-6]|div|table|ul|ol)\b', letter_text, re.IGNORECASE))

    return build_docx_response(
        content=letter_text,
        title=f"Expert Letter for {client_name}",
        filename_stem=f"Expert_Letter_{client_clean}_signed_by_{expert_clean}",
        doc_type="Expert Letter" if language == 'en' else "Carta de Experto",
        author=expert_name,
        language=language,
        is_html=is_html,
        add_cover=False,
    )


@router.delete("/{letter_id}")
async def delete_expert_letter(letter_id: str, current_user=Depends(get_current_user_wrapper)):
    """Soft delete expert letter."""
    db = get_db()
    try:
        letter = await db.expert_letters.find_one({"id": letter_id}, {"_id": 0})
        if not letter:
            raise HTTPException(status_code=404, detail="Letter not found")

        result = await db.expert_letters.update_one(
            {"id": letter_id},
            {"$set": {
                "status": "deleted",
                "deleted_at": datetime.now(timezone.utc).isoformat(),
                "deleted_by": current_user.id
            }}
        )
        if result.modified_count == 0:
            raise HTTPException(status_code=400, detail="Failed to delete letter")

        return {"message": "Carta de experto movida a la papelera"}
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error deleting expert letter: {str(e)}")
        raise HTTPException(status_code=500, detail=str(e))
